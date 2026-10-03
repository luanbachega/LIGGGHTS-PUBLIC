import sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import model as M
OUT = sys.argv[1]
wb = Workbook(); ws = wb.active; ws.title = 'Premissas'
B = Font(bold=True); IN = PatternFill('solid', fgColor='FFF2CC')
prem = [('Preço 2Q (R$/m² privativo)', M.P['preco_2q']), ('Preço studio (R$/m² privativo)', M.P['preco_st']),
        ('Preço vaga excedente (R$)', M.P['preco_vaga']), ('Custo por unidade adicional (R$)', M.P['custo_un_extra']),
        ('CUB-PR ago/2026 (R$/m²)', M.P['cub_pr']), ('GI estimado EIV', M.P['gi_eiv'])]
ws['A1'] = '151 Estudo de Viabilidade – premissas (células amarelas são editáveis)'; ws['A1'].font = B
for i, (k, v) in enumerate(prem, start=3):
    ws.cell(i, 1, k); c = ws.cell(i, 2, v); c.fill = IN; c.number_format = '#,##0.00'
ws.column_dimensions['A'].width = 38; ws.column_dimensions['B'].width = 16
e = wb.create_sheet('Estudos')
hdr = ['Estudo', 'Unid.', 'Aptos 2Q', 'Studios', 'Priv. 2Q (m²)', 'Priv. studios (m²)', 'Vagas', 'Vagas exigidas',
       'Excedentes', 'Área construída (m²)', 'Custo fixo (R$)', 'Unid. adicionais vs R00', 'Custo obra (R$)',
       'VGV (R$)', 'AP (EIV)', 'Exige EIV (1/0)', 'Contrapartida EIV (R$)', 'Resultado (R$)', 'Margem']
for j, h in enumerate(hdr, 1):
    c = e.cell(1, j, h); c.font = B; c.alignment = Alignment(wrap_text=True, vertical='top')
base = {'A2': 'A', 'B2': 'B', 'C2': 'C', 'C3': 'C', 'D2': 'D', 'A2-100': 'A', 'B2-100': 'B', 'D2-100': 'D'}
BY = {s['key']: s for s in M.ORIG + M.NEW + M.V100}
for i, s in enumerate(M.ORIG + M.NEW + M.V100, start=2):
    p2 = sum(a * q for a, q in s['units'].items() if not M.is_studio(a))
    ps = sum(a * q for a, q in s['units'].items() if M.is_studio(a))
    dun = s['nun'] - BY[base[s['key']]]['nun'] if s['key'] in base else 0
    fixed = s['custo'] - dun * M.P['custo_un_extra']
    vals = [s['curto'], s['nun'], s['n2q'], s['nst'], round(p2, 2), round(ps, 2), s['vagas'], s['vx']['total'],
            f'=MAX(0,G{i}-H{i})', round(s['constr'], 2), round(fixed, 2), dun,
            f'=K{i}+L{i}*Premissas!$B$6', f'=E{i}*Premissas!$B$3+F{i}*Premissas!$B$4+I{i}*Premissas!$B$5',
            1.25 if s['nun'] > 200 else 1, 1 if s['nun'] >= 130 else 0,
            f'=P{i}*Premissas!$B$8*O{i}*0.25*J{i}*Premissas!$B$7', f'=N{i}-M{i}-Q{i}', f'=R{i}/N{i}']
    for j, v in enumerate(vals, 1):
        c = e.cell(i, j, v)
        if j in (5, 6, 10, 11, 13, 14, 17, 18): c.number_format = '#,##0'
        if j == 19: c.number_format = '0%'
e.column_dimensions['A'].width = 34
for col in 'BCDEFGHIJKLMNOPQRS': e.column_dimensions[col].width = 14
e.freeze_panes = 'B2'
n = e.max_row + 2
e.cell(n, 1, 'Custo fixo = custo da obra do R00 ± diferenças de área/garagem (ver PDF comparativo); a parcela por unidade adicional é recalculada pela premissa.').font = Font(italic=True, size=9)
e.cell(n + 1, 1, 'Linhas -100 são indicativas (sem prancha). Vagas exigidas: LC 26/2020 art. 259 + Anexo VII (studios 1 vaga a cada 3; 2Q 1/un.; visitantes 5%).').font = Font(italic=True, size=9)
fn = os.path.join(OUT, '151 MODELO NUMERICO - VERSOES COM STUDIOS.xlsx'); wb.save(fn); print('ok', fn)
