import sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import model153 as M
wb = Workbook(); ws = wb.active; ws.title = 'Premissas'
B = Font(bold=True); IN = PatternFill('solid', fgColor='FFF2CC')
prem = [('Preço 2Q (R$/m² privativo)', M.P['preco_2q']), ('Preço 3Q (R$/m² privativo)', M.P['preco_3q']),
        ('Preço studio (R$/m² privativo)', M.P['preco_st']), ('Preço vaga excedente (R$)', M.P['preco_vaga']),
        ('Custo R$/m² construído (torre/térreo/ático)', M.P['custo_torre']), ('Custo R$/m² subsolo', M.P['custo_subsolo']),
        ('Custo extra por studio (R$)', M.P['custo_st_extra']), ('Demolição (R$)', M.P['demolicao']),
        ('Terreno – valor venal CND (R$)', 3282771.82)]
ws['A1'] = '153 Estudo de Viabilidade – premissas (células amarelas são editáveis)'; ws['A1'].font = B
for i, (k, v) in enumerate(prem, start=3):
    ws.cell(i, 1, k); c = ws.cell(i, 2, v); c.fill = IN; c.number_format = '#,##0.00'
ws.column_dimensions['A'].width = 44; ws.column_dimensions['B'].width = 16
e = wb.create_sheet('Opcoes')
hdr = ['Opção', 'Unid.', '2Q', '3Q', 'Studios', 'Priv. 2Q (m²)', 'Priv. 3Q (m²)', 'Priv. studios (m²)', 'Vagas', 'Exigidas',
       'Excedentes', 'Construída exceto subsolo (m²)', 'Subsolo (m²)', 'Custo obra (R$)', 'VGV (R$)', 'Resultado (R$)',
       'Margem', 'Resultado − terreno (R$)']
for j, h in enumerate(hdr, 1):
    c = e.cell(1, j, h); c.font = B; c.alignment = Alignment(wrap_text=True, vertical='top')
for i, o in enumerate(M.OPTS, start=2):
    pa = lambda t: round(sum(a * q for (tt, a), q in o['units'].items() if tt == t), 2)
    vals = [o['label'], o['nun'], o['n2'], o['n3'], o['nst'], pa('2Q'), pa('3Q'), pa('ST'), o['vagas'], o['vx']['total'],
            f'=MAX(0,I{i}-J{i})', round(o['constr'] - o['sub_area'], 2), round(o['sub_area'], 2),
            f'=L{i}*Premissas!$B$7+M{i}*Premissas!$B$8+E{i}*Premissas!$B$9+Premissas!$B$10',
            f'=F{i}*Premissas!$B$3+G{i}*Premissas!$B$4+H{i}*Premissas!$B$5+K{i}*Premissas!$B$6',
            f'=O{i}-N{i}', f'=P{i}/O{i}', f'=P{i}-Premissas!$B$11']
    for j, v in enumerate(vals, 1):
        c = e.cell(i, j, v)
        if j >= 6 and j != 17: c.number_format = '#,##0'
        if j == 17: c.number_format = '0%'
e.column_dimensions['A'].width = 40
for col in 'BCDEFGHIJKLMNOPQR': e.column_dimensions[col].width = 14
n = e.max_row + 2
e.cell(n, 1, 'Nenhuma opção atinge 130 unidades: sem contrapartida de EIV. Custos ±20% (CUB-PR ago/2026, R8-N).').font = Font(italic=True, size=9)
fn = os.path.join(sys.argv[1], '153 MODELO NUMERICO - OPCOES.xlsx'); wb.save(fn); print('ok', fn)
